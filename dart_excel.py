"""
재무제표 엑셀(xlsx) 생성

  - 시트 순서: 정보 → 재무상태표 → (포괄)손익계산서 → 자본변동표 → 현금흐름표 → 이익잉여금처분계산서
  - 금액은 원 단위로 통일, 천 단위 구분기호, 음수는 괄호, 0은 '-'
  - 머리글 행은 가운데 정렬·굵게
"""

import io
from dataclasses import dataclass, field
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from dart_parser import STATEMENT_ORDER, ParseResult, Statement, sheet_name_for

AMOUNT_FORMAT = '#,##0;(#,##0);"-"'
DECIMAL_FORMAT = '#,##0.##;(#,##0.##);"-"'
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
HEADER_FONT = Font(bold=True)
TITLE_FONT = Font(bold=True, size=14)
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

XBRL_SHEETS = [
    ("BS", "재무상태표"),
    ("IS", "손익계산서"),
    ("CIS", "포괄손익계산서"),
    ("SCE", "자본변동표"),
    ("CF", "현금흐름표"),
]


def _info_sheet(wb: Workbook, meta: dict, notes: list[str]) -> None:
    ws = wb.active
    ws.title = "정보"
    ws["A1"] = "DART 재무제표 추출"
    ws["A1"].font = TITLE_FONT
    rows = [
        ("회사명", meta.get("corp_name", "")),
        ("고유번호", meta.get("corp_code", "")),
        ("종목코드", meta.get("stock_code", "") or "비상장"),
        ("보고서명", meta.get("report_nm", "")),
        ("접수번호", meta.get("rcept_no", "")),
        ("접수일자", meta.get("rcept_dt", "")),
        ("추출방식", meta.get("method", "")),
        ("DART 원문", f"https://dart.fss.or.kr/dsaf001/main.do?rcpNo={meta.get('rcept_no', '')}"),
        ("생성일시", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]
    for i, (k, v) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=k).font = HEADER_FONT
        cell = ws.cell(row=i, column=2, value=v)
        if k == "DART 원문":
            cell.hyperlink = v
            cell.style = "Hyperlink"

    r = len(rows) + 4
    ws.cell(row=r, column=1, value="확인사항").font = HEADER_FONT
    for note in notes:
        r += 1
        ws.cell(row=r, column=1, value="•")
        ws.cell(row=r, column=2, value=note)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 90


def _write_grid(ws: Worksheet, rows: list[list], header_rows: int, start_row: int) -> None:
    width = max((len(r) for r in rows), default=0)
    for r_idx, row in enumerate(rows):
        excel_row = start_row + r_idx
        for c_idx in range(width):
            value = row[c_idx] if c_idx < len(row) else ""
            cell = ws.cell(row=excel_row, column=c_idx + 1, value=value if value != "" else None)
            cell.border = BORDER
            if r_idx < header_rows:
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL
                cell.alignment = CENTER
            elif isinstance(value, (int, float)):
                cell.number_format = DECIMAL_FORMAT if isinstance(value, float) else AMOUNT_FORMAT

    ws.column_dimensions["A"].width = 42
    for c in range(2, width + 1):
        ws.column_dimensions[get_column_letter(c)].width = 20
    ws.freeze_panes = ws.cell(row=start_row + header_rows, column=2)


def _write_statement_heading(ws: Worksheet, st: Statement) -> int:
    """제목·기간·단위 문구를 쓰고 마지막으로 쓴 행 번호 반환"""
    ws["A1"] = st.title
    ws["A1"].font = TITLE_FONT
    line = 2
    for text in st.period_texts[:3]:
        ws.cell(row=line, column=1, value=text)
        line += 1
    unit_note = "(단위: 원)" if st.unit == "원" else f"(단위: 원 / 원문 단위 {st.unit}를 원으로 환산)"
    ws.cell(row=line, column=1, value=unit_note)
    return line


def build_workbook_from_document(meta: dict, result: ParseResult) -> bytes:
    wb = Workbook()
    notes = [
        "금액은 모두 원 단위로 환산하였습니다. (주당이익 행은 원문 그대로 원 단위)",
        "원문 표를 자동으로 읽은 결과이므로, 합계 검증 및 DART 원문과의 대사를 권장합니다.",
    ]

    found = [st for cons in (True, False) for k in STATEMENT_ORDER if (st := result.get(k, cons))]
    for st in found:
        ws = wb.create_sheet(sheet_name_for(st.key, st.title))
        line = _write_statement_heading(ws, st)
        _write_grid(ws, st.rows, st.header_row_count, start_row=line + 2)
        if not st.unit_found:
            notes.append(f"{st.title}: 원문에서 단위 표시를 찾지 못해 '원'으로 가정했습니다. 원문 확인 필요.")

    for cons in sorted({st.consolidated for st in found}, reverse=True):
        missing = [name for k, name in (("BS", "재무상태표"), ("SCE", "자본변동표"), ("CF", "현금흐름표"))
                   if not result.get(k, cons)]
        if not result.income_statement(cons):
            missing.insert(1, "손익계산서")
        if missing:
            prefix = "연결" if cons else ""
            notes.append("찾지 못한 재무제표: " + ", ".join(prefix + m for m in missing))

    if not found:
        notes.append("재무제표 표를 식별하지 못해 원문의 숫자 표를 '원문표' 시트에 그대로 실었습니다.")
        for i, grid in enumerate(result.raw_tables[:30], start=1):
            ws = wb.create_sheet(f"원문표{i}")
            _write_grid(ws, grid, 0, start_row=1)

    _info_sheet(wb, {**meta, "method": "공시 원문(감사보고서) 표 파싱"}, notes)
    return _to_bytes(wb)


@dataclass
class SheetTable:
    """시트 하나에 쓸 표 (원문 파싱·XBRL 공통)"""
    title: str
    rows: list[list]
    header_rows: int
    unit_note: str = "(단위: 원)"
    period_texts: list[str] = field(default_factory=list)


def table_from_statement(st: Statement) -> SheetTable:
    unit_note = "(단위: 원)" if st.unit == "원" else f"(단위: 원 / 원문 단위 {st.unit}를 원으로 환산)"
    return SheetTable(st.title, st.rows, st.header_row_count, unit_note, st.period_texts[:3])


def table_from_xbrl(items: list[dict], sj_div: str, fs_div: str) -> SheetTable | None:
    part = sorted((it for it in items if it.get("sj_div") == sj_div), key=lambda it: int(it.get("ord") or 0))
    if not part:
        return None
    first = part[0]
    headers = ["계정명", "계정ID"]
    if sj_div == "SCE":
        headers.append("자본구성요소")
    period_names = [first.get("thstrm_nm") or "당기", first.get("frmtrm_nm") or "전기"]
    if first.get("bfefrmtrm_nm"):
        period_names.append(first["bfefrmtrm_nm"])
    headers += period_names

    rows: list[list] = [headers]
    for it in part:
        row = [it.get("account_nm", ""), it.get("account_id", "")]
        if sj_div == "SCE":
            row.append(it.get("account_detail", ""))
        amounts = [it.get("thstrm_amount"), it.get("frmtrm_amount"), it.get("bfefrmtrm_amount")]
        row += [_xbrl_amount(a) for a in amounts[: len(period_names)]]
        rows.append(row)
    name = dict(XBRL_SHEETS)[sj_div]
    return SheetTable(("연결" if fs_div == "CFS" else "") + name, rows, 1)


def xbrl_income_statement(items: list[dict], fs_div: str) -> SheetTable | None:
    """손익계산서가 없으면(단일 포괄손익계산서 양식) 포괄손익계산서"""
    return table_from_xbrl(items, "IS", fs_div) or table_from_xbrl(items, "CIS", fs_div)


def _write_table(ws: Worksheet, table: SheetTable) -> None:
    ws["A1"] = table.title
    ws["A1"].font = TITLE_FONT
    line = 2
    for text in table.period_texts:
        ws.cell(row=line, column=1, value=text)
        line += 1
    ws.cell(row=line, column=1, value=table.unit_note)
    _write_grid(ws, table.rows, table.header_rows, start_row=line + 2)
    if table.rows and table.rows[0][:2] == ["계정명", "계정ID"]:
        ws.column_dimensions["B"].width = 30


def build_workbook_from_xbrl(meta: dict, items: list[dict], bsns_year: int, fs_div: str) -> bytes:
    """fnlttSinglAcntAll 응답을 재무제표별 시트로"""
    wb = Workbook()
    fs_label = "연결" if fs_div == "CFS" else "별도"
    notes = [f"OpenDART XBRL 재무제표({fs_label}, 사업연도 {bsns_year}) 기준이며 금액 단위는 원입니다."]
    for sj_div, name in XBRL_SHEETS:
        table = table_from_xbrl(items, sj_div, fs_div)
        if table:
            _write_table(wb.create_sheet(f"{fs_label}{name}"), table)
    _info_sheet(wb, {**meta, "method": f"OpenDART XBRL 재무제표 ({fs_label})"}, notes)
    return _to_bytes(wb)


@dataclass
class YearResult:
    year: int
    report_nm: str = ""
    rcept_no: str = ""
    method: str = ""
    bs: SheetTable | None = None
    pl: SheetTable | None = None
    notes: list[str] = field(default_factory=list)


def build_annual_workbook(meta: dict, results: list[YearResult], fs_div: str) -> bytes:
    """연도별 BS(2025), PL(2025) 시트. 첫 시트는 연도별 출처 요약"""
    wb = Workbook()
    ws = wb.active
    ws.title = "정보"
    fs_label = "연결" if fs_div == "CFS" else "별도"
    ws["A1"] = f"{meta.get('corp_name', '')} 연도별 재무제표 ({fs_label} 우선)"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (f"고유번호 {meta.get('corp_code', '')} · 종목코드 {meta.get('stock_code') or '비상장'} · "
                f"생성일시 {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    ws["A3"] = "금액은 원 단위입니다. 원문 파싱 결과는 합계 검증 및 DART 원문과의 대사를 권장합니다."

    headers = ["연도", "보고서", "접수번호", "추출방식", "시트", "확인사항"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=c, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = HEADER_FONT, HEADER_FILL, CENTER, BORDER

    for i, res in enumerate(sorted(results, key=lambda r: r.year, reverse=True), start=6):
        sheets = []
        for prefix, table in (("BS", res.bs), ("PL", res.pl)):
            if table:
                name = f"{prefix}({res.year})"
                _write_table(wb.create_sheet(name), table)
                sheets.append(name)
        values = [res.year, res.report_nm or "해당 연도 보고서 없음", res.rcept_no, res.method,
                  ", ".join(sheets), " / ".join(res.notes)]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if res.rcept_no:
            link = ws.cell(row=i, column=3)
            link.hyperlink = f"https://dart.fss.or.kr/dsaf001/main.do?rcpNo={res.rcept_no}"
            link.style = "Hyperlink"
            link.border = BORDER

    for col, width in zip("ABCDEF", (8, 30, 18, 26, 22, 60)):
        ws.column_dimensions[col].width = width
    return _to_bytes(wb)


def _xbrl_amount(value: str | None):
    if value is None:
        return ""
    v = str(value).replace(",", "").strip()
    if v in ("", "-"):
        return ""
    try:
        return int(v)
    except ValueError:
        try:
            return float(v)
        except ValueError:
            return value


def _to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
