import asyncio
import io
import zipfile
from pathlib import Path

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import dart_client
from dart_excel import build_workbook_from_document, build_workbook_from_xbrl
from dart_parser import detect_statement, detect_unit, parse_document, parse_number
from kakao_tax_skill_server import app

SAMPLE = (Path(__file__).parent / "sample_audit_report.xml").read_text(encoding="utf-8")


@pytest.mark.parametrize("text, expected", [
    ("1,234,567", 1234567),
    ("(1,234)", -1234),
    ("△50,000", -50000),
    ("-1,000", -1000),
    ("-", 0),
    ("12.5", 12.5),
    ("4,5", None),          # 주석번호 '4,5'는 천단위 형식이 아님
    ("2024.12.31", None),
    ("제 10(당) 기", None),
    ("(1,234", None),
])
def test_parse_number(text, expected):
    assert parse_number(text) == expected


def test_detect_statement_ignores_long_sentences():
    assert detect_statement("재 무 상 태 표") == ("BS", "재무상태표")
    assert detect_statement("연 결 포 괄 손 익 계 산 서") == ("CIS", "연결포괄손익계산서")
    assert detect_statement("손익계산서") == ("IS", "손익계산서")
    assert detect_statement("2024년 12월 31일 현재의 재무상태표, 동일로 종료되는 보고기간의 손익계산서") is None
    assert detect_statement("2-1. 연결 재무상태표") == ("BS", "연결재무상태표")
    assert detect_statement("(2) 요약 재무상태표") is None
    assert detect_statement("1) 재무상태표") is None
    assert detect_statement("관계기업 재무상태표") is None
    assert detect_unit("(단위 : 천원)") == "천원"
    assert detect_unit("(단위: 백만원)") == "백만원"


def test_parse_audit_report():
    result = parse_document([SAMPLE])
    assert set(result.statements) == {"BS", "IS", "CF"}

    bs = result.statements["BS"]
    assert bs.unit == "천원" and bs.header_row_count == 2
    rows = {r[0]: r for r in bs.rows}
    assert rows["현금및현금성자산"][1] == "4,5"                 # 주석 열은 문자 그대로
    assert rows["현금및현금성자산"][2] == 300_000_000          # 천원 -> 원
    assert rows["대손충당금"][2] == -50_000_000
    assert rows["대손충당금"][4] == -50_000_000                # △ 표기
    assert rows["자산총계"][3] == 1_500_000_000

    is_rows = {r[0]: r for r in result.statements["IS"].rows}
    assert is_rows["Ⅲ. 당기순이익(손실)"][3] == -30_000_000
    assert is_rows["Ⅳ. 주당이익(손실)"][2] == 1_200            # 주당이익은 환산하지 않음

    cf = result.statements["CF"]
    assert cf.unit == "원" and cf.rows[1][1] == 10_000_000


def test_notes_table_does_not_override_statement():
    result = parse_document([SAMPLE])
    values = [v for r in result.statements["BS"].rows for v in r]
    assert 999_999_000 not in values and 999_999 not in values


def test_document_workbook():
    result = parse_document([SAMPLE])
    content = build_workbook_from_document({"corp_name": "테스트건설", "rcept_no": "20250320000123"}, result)
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["정보", "재무상태표", "손익계산서", "현금흐름표"]
    ws = wb["재무상태표"]
    header = next(c for c in ws["A"] if c.value == "과 목")
    assert header.alignment.horizontal == "center"
    amount = next(row for row in ws.iter_rows() if row[0].value == "매출채권")[2]
    assert amount.value == 1_250_000_000 and amount.number_format.startswith("#,##0")
    notes = [c.value for c in wb["정보"]["B"] if c.value]
    assert any("자본변동표" in str(n) for n in notes)     # 누락 재무제표 안내


def test_document_workbook_falls_back_to_raw_tables():
    html = "<table>" + "".join(f"<tr><td>항목{i}</td><td>{i},000</td><td>{i},500</td></tr>" for i in range(5)) + "</table>"
    wb = load_workbook(io.BytesIO(build_workbook_from_document({}, parse_document([html]))))
    assert "원문표1" in wb.sheetnames


def test_xbrl_workbook():
    items = [
        {"sj_div": "BS", "ord": "2", "account_id": "ifrs-full_Assets", "account_nm": "자산총계",
         "thstrm_nm": "제 56 기", "thstrm_amount": "455905980000000", "frmtrm_nm": "제 55 기",
         "frmtrm_amount": "448424507000000", "bfefrmtrm_nm": "제 54 기", "bfefrmtrm_amount": "426621158000000"},
        {"sj_div": "BS", "ord": "1", "account_id": "ifrs-full_CurrentAssets", "account_nm": "유동자산",
         "thstrm_nm": "제 56 기", "thstrm_amount": "195936557000000", "frmtrm_nm": "제 55 기",
         "frmtrm_amount": "218470581000000", "bfefrmtrm_nm": "제 54 기", "bfefrmtrm_amount": "-"},
    ]
    wb = load_workbook(io.BytesIO(build_workbook_from_xbrl({}, items, 2024, "CFS")))
    ws = wb["연결재무상태표"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=4)]
    assert rows[0] == ["계정명", "계정ID", "제 56 기", "제 55 기", "제 54 기"]
    assert rows[1][0] == "유동자산" and rows[1][4] is None      # ord 순 정렬, '-'는 빈칸
    assert rows[2][2] == 455_905_980_000_000


def test_search_corps_ranking():
    xml = """<result>
      <list><corp_code>00000001</corp_code><corp_name>테스트건설산업</corp_name><stock_code> </stock_code><modify_date>20240101</modify_date></list>
      <list><corp_code>00000002</corp_code><corp_name>(주)테스트건설</corp_name><stock_code></stock_code><modify_date>20240101</modify_date></list>
      <list><corp_code>00000003</corp_code><corp_name>대한테스트건설</corp_name><stock_code>123450</stock_code><modify_date>20240101</modify_date></list>
    </result>""".encode("utf-8")
    corps = dart_client.parse_corp_code_xml(xml)
    hits = dart_client.search_corps(corps, "테스트 건설")
    assert [h["corp_code"] for h in hits] == ["00000002", "00000001", "00000003"]
    assert "norm" not in hits[0]


def test_year_from_report_name():
    assert dart_client.year_from_report_name("사업보고서 (2024.12)") == 2024
    assert dart_client.year_from_report_name("감사보고서 (2024.12)") == 2024
    assert dart_client.year_from_report_name("감사보고서") is None


def test_excel_endpoint(monkeypatch):
    async def fake_fetch(client, rcept_no):
        return [(f"{rcept_no}.xml", SAMPLE)]
    monkeypatch.setattr(dart_client, "fetch_document_files", fake_fetch)

    client = TestClient(app)
    assert "DART 재무제표 조회" in client.get("/dart").text

    resp = client.get("/dart/api/excel", params={
        "rcept_no": "20250320000123", "corp_code": "00123456", "corp_name": "테스트건설",
        "report_nm": "감사보고서 (2024.12)", "rcept_dt": "20250320", "mode": "doc"})
    assert resp.status_code == 200
    assert "UTF-8''" in resp.headers["content-disposition"]
    assert load_workbook(io.BytesIO(resp.content)).sheetnames[1] == "재무상태표"

    bad = client.get("/dart/api/excel", params={"rcept_no": "../etc", "corp_code": "00123456"})
    assert bad.status_code == 422


def test_dart_error_is_reported(monkeypatch):
    async def failing(client, rcept_no):
        raise dart_client.DartError("OpenDART 오류 [020] 요청 제한을 초과하였습니다.")
    monkeypatch.setattr(dart_client, "fetch_document_files", failing)
    resp = TestClient(app).get("/dart/api/excel", params={"rcept_no": "20250320000123", "corp_code": "00123456"})
    assert resp.status_code == 502 and "020" in resp.json()["detail"]


def test_fetch_document_files_reads_zip(monkeypatch):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("20250320000123_00760.xml", "<p>첨부</p>")
        zf.writestr("20250320000123.xml", SAMPLE.encode("utf-8"))
        zf.writestr("20250320000123_0.xml", '<?xml version="1.0" encoding="euc-kr"?><p>재무상태표</p>'.encode("cp949"))
    monkeypatch.setattr(dart_client, "DART_API_KEY", "test")

    async def run():
        transport = httpx.MockTransport(lambda request: httpx.Response(200, content=buf.getvalue()))
        async with httpx.AsyncClient(transport=transport) as c:
            return await dart_client.fetch_document_files(c, "20250320000123")

    files = dict(asyncio.run(run()))
    assert list(files)[0] == "20250320000123.xml"                       # 본문 파일 우선
    assert "재무상태표" in files["20250320000123_0.xml"]                   # EUC-KR 원문 디코딩


def test_document_error_body_is_reported(monkeypatch):
    monkeypatch.setattr(dart_client, "DART_API_KEY", "test")
    body = "<result><status>014</status><message>파일이 존재하지 않습니다.</message></result>".encode("utf-8")

    async def run():
        transport = httpx.MockTransport(lambda request: httpx.Response(200, content=body))
        async with httpx.AsyncClient(transport=transport) as c:
            await dart_client.fetch_document_files(c, "20250320000123")

    with pytest.raises(dart_client.DartError, match="014"):
        asyncio.run(run())


# ---------------------------------------------------------------------------
# 연도별 BS·PL 엑셀 (/dart/api/annual)
# ---------------------------------------------------------------------------

CONSOLIDATED = (SAMPLE.replace("재 무 상 태 표", "연 결 재 무 상 태 표")
                .replace("손 익 계 산 서", "연 결 손 익 계 산 서")
                .replace("현 금 흐 름 표", "연 결 현 금 흐 름 표")
                .replace("<TE>300,000</TE>", "<TE>777,000</TE>"))

XBRL_ITEMS = [
    {"sj_div": "BS", "ord": "1", "account_id": "ifrs-full_Assets", "account_nm": "자산총계",
     "thstrm_nm": "제 56 기", "thstrm_amount": "1000", "frmtrm_nm": "제 55 기", "frmtrm_amount": "900"},
    {"sj_div": "CIS", "ord": "1", "account_id": "ifrs-full_Revenue", "account_nm": "수익(매출액)",
     "thstrm_nm": "제 56 기", "thstrm_amount": "500", "frmtrm_nm": "제 55 기", "frmtrm_amount": "400"},
]


def _report(rcept_no, report_nm):
    return {"rcept_no": rcept_no, "report_nm": report_nm, "rcept_dt": rcept_no[:8], "flr_nm": "",
            "fiscal_year": dart_client.year_from_report_name(report_nm),
            "is_annual": "사업보고서" in report_nm, "is_consolidated": "연결" in report_nm}


def _patch_dart(monkeypatch, reports, documents, xbrl=None):
    async def fake_list(client, corp_code, years=10, begin=None):
        return reports

    async def fake_fetch(client, rcept_no):
        return [(f"{rcept_no}.xml", text) for text in documents[rcept_no]]

    async def fake_xbrl(client, corp_code, year, fs_div, try_previous=True):
        if xbrl is None or year not in xbrl:
            raise dart_client.DartError("XBRL 재무제표 데이터가 없습니다.")
        return year, xbrl[year]

    monkeypatch.setattr(dart_client, "list_reports", fake_list)
    monkeypatch.setattr(dart_client, "fetch_document_files", fake_fetch)
    monkeypatch.setattr(dart_client, "fetch_xbrl_statements", fake_xbrl)


def _annual(params):
    resp = TestClient(app).get("/dart/api/annual", params={"corp_code": "00123456", "corp_name": "테스트건설", **params})
    assert resp.status_code == 200, resp.text
    return load_workbook(io.BytesIO(resp.content))


def _value(ws, account, col):
    return next(row for row in ws.iter_rows(values_only=True) if row[0] == account)[col]


def test_annual_unlisted_by_year(monkeypatch):
    reports = [_report("20250320000002", "감사보고서 (2024.12)"), _report("20240320000001", "감사보고서 (2023.12)")]
    _patch_dart(monkeypatch, reports, {"20250320000002": [SAMPLE], "20240320000001": [SAMPLE]})
    wb = _annual({"years": "2022,2023,2024"})
    assert wb.sheetnames == ["정보", "BS(2024)", "PL(2024)", "BS(2023)", "PL(2023)"]
    assert _value(wb["BS(2024)"], "매출채권", 2) == 1_250_000_000
    info = [row for row in wb["정보"].iter_rows(min_row=6, values_only=True)]
    assert [r[0] for r in info] == [2024, 2023, 2022]
    assert info[2][1] == "해당 연도 보고서 없음"


def test_annual_selects_separate_or_consolidated(monkeypatch):
    reports = [_report("20250320000002", "감사보고서 (2024.12)"), _report("20250320000003", "연결감사보고서 (2024.12)")]
    docs = {"20250320000002": [SAMPLE], "20250320000003": [CONSOLIDATED]}
    _patch_dart(monkeypatch, reports, docs)
    assert _value(_annual({"years": "2024", "fs_div": "OFS"})["BS(2024)"], "현금및현금성자산", 2) == 300_000_000
    wb = _annual({"years": "2024", "fs_div": "CFS"})
    assert _value(wb["BS(2024)"], "현금및현금성자산", 2) == 777_000_000
    assert wb["BS(2024)"]["A1"].value == "연결재무상태표"


def test_annual_falls_back_when_consolidated_missing(monkeypatch):
    _patch_dart(monkeypatch, [_report("20250320000002", "감사보고서 (2024.12)")], {"20250320000002": [SAMPLE]})
    wb = _annual({"years": "2024", "fs_div": "CFS"})
    assert "BS(2024)" in wb.sheetnames
    notes = next(wb["정보"].iter_rows(min_row=6, values_only=True))[5]
    assert "별도 재무제표를 실었습니다" in notes


def test_annual_listed_uses_xbrl_then_document(monkeypatch):
    reports = [_report("20250311000010", "사업보고서 (2024.12)"), _report("20240311000009", "사업보고서 (2023.12)")]
    # 2024는 XBRL, 2023은 XBRL 실패 -> 사업보고서 원문(연결+별도 수록) 파싱
    docs = {"20240311000009": [CONSOLIDATED, SAMPLE]}
    _patch_dart(monkeypatch, reports, docs, xbrl={2024: XBRL_ITEMS})
    wb = _annual({"years": "2023,2024", "fs_div": "OFS"})
    assert wb.sheetnames == ["정보", "BS(2024)", "PL(2024)", "BS(2023)", "PL(2023)"]
    assert _value(wb["PL(2024)"], "수익(매출액)", 2) == 500                 # IS 없으면 CIS
    assert _value(wb["BS(2023)"], "현금및현금성자산", 2) == 300_000_000      # 별도 선택
    info = list(wb["정보"].iter_rows(min_row=6, values_only=True))
    assert info[0][3] == "XBRL (별도)" and info[1][3] == "원문 파싱"


def test_annual_rejects_bad_years():
    client = TestClient(app)
    assert client.get("/dart/api/annual", params={"corp_code": "00123456", "years": "abcd"}).status_code == 422
    assert client.get("/dart/api/annual", params={"corp_code": "00123456", "years": "2099"}).status_code == 400


# ---------------------------------------------------------------------------
# 엑셀 VBA용: 회사명만으로 조회, 접근 토큰
# ---------------------------------------------------------------------------

CORPS = [
    {"corp_code": "00123456", "corp_name": "테스트건설", "stock_code": "", "modify_date": "", "norm": "테스트건설"},
    {"corp_code": "00123457", "corp_name": "테스트건설산업", "stock_code": "", "modify_date": "", "norm": "테스트건설산업"},
    {"corp_code": "00200001", "corp_name": "동명상사", "stock_code": "", "modify_date": "", "norm": "동명상사"},
    {"corp_code": "00200002", "corp_name": "동명상사", "stock_code": "012340", "modify_date": "", "norm": "동명상사"},
]


def test_resolve_corp():
    assert dart_client.resolve_corp(CORPS, "(주)테스트건설")[0]["corp_code"] == "00123456"   # 정확 일치 우선
    assert dart_client.resolve_corp(CORPS, "건설산업")[0]["corp_code"] == "00123457"        # 부분일치 1건
    corp, candidates = dart_client.resolve_corp(CORPS, "동명상사")
    assert corp is None and len(candidates) == 2
    assert dart_client.resolve_corp(CORPS, "없는회사") == (None, [])


def test_annual_by_name(monkeypatch):
    async def fake_corps(client):
        return CORPS
    monkeypatch.setattr(dart_client, "load_corp_list", fake_corps)
    _patch_dart(monkeypatch, [_report("20250320000002", "감사보고서 (2024.12)")], {"20250320000002": [SAMPLE]})
    client = TestClient(app)

    resp = client.get("/dart/api/annual", params={"corp_name": "테스트건설", "years": "2024"})
    assert resp.status_code == 200
    assert "테스트건설" in load_workbook(io.BytesIO(resp.content))["정보"]["A1"].value

    dup = client.get("/dart/api/annual", params={"corp_name": "동명상사", "years": "2024"})
    assert dup.status_code == 409
    assert dup.text.splitlines() == ["00200002\t동명상사\t상장 012340", "00200001\t동명상사\t비상장"]  # 상장사 우선

    assert client.get("/dart/api/annual", params={"corp_name": "없는회사", "years": "2024"}).status_code == 404
    assert client.get("/dart/api/annual", params={"years": "2024"}).status_code == 400


def test_access_token(monkeypatch):
    monkeypatch.setenv("DART_WEB_TOKEN", "secret")
    client = TestClient(app)
    assert client.get("/dart/api/annual", params={"years": "2024"}).status_code == 401
    assert client.get("/dart/api/annual", params={"years": "2024", "token": "wrong"}).status_code == 401
    assert client.get("/dart/api/annual", params={"years": "2024"}, headers={"X-Access-Token": "secret"}).status_code == 400
    assert client.get("/dart").status_code == 200     # 페이지 자체는 공개, API만 보호


def test_vba_import_file_matches_source():
    """excel/DartFS.bas(가져오기용, CP949·CRLF)는 DartFS_utf8.bas와 내용이 같아야 함"""
    excel_dir = Path(__file__).parent.parent / "excel"
    source = (excel_dir / "DartFS_utf8.bas").read_text(encoding="utf-8").replace("\r\n", "\n")
    imported = (excel_dir / "DartFS.bas").read_bytes()
    assert imported == source.replace("\n", "\r\n").encode("cp949")
    assert imported.startswith(b'Attribute VB_Name = "DartFS"')
